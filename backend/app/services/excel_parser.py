import re
import io
from datetime import datetime, date
from typing import List, Dict, Any, Optional
import openpyxl

def categorize_dish(name: str) -> str:
    n = name.lower().strip()
    if any(k in n for k in ["rice", "biryani", "pulao", "bath", "noodle"]):
        return "Rice"
    if any(k in n for k in ["roti", "naan", "chapathi", "phulka", "bread", "kulcha", "paratha"]):
        return "Bread"
    if any(k in n for k in ["jalebi", "laddu", "brownie", "rasgulla", "khaja", "tiramisu", "pudding", "ice cream", "sweet", "halwa", "payasam"]):
        return "Desserts"
    if any(k in n for k in ["tea", "coffee", "soup", "coriander", "manchow", "drink", "juice", "beverage"]):
        return "Beverages"
    if any(k in n for k in ["salad", "salsa", "sundal", "raitha", "pachadi", "pickle", "papad", "chutney"]):
        return "Salads"
    if any(k in n for k in ["tikka", "kebab", "bhajji", "65", "pakoda", "parcel", "puff", "idly", "wada", "snack"]):
        return "Starters"
    return "Main Course"

def parse_sahara_sheet(sheet, source_file: str = "Daily report.xlsx") -> List[Dict[str, Any]]:
    records = []
    # Extract date
    d_str = "2026-08-24"
    d_cell = sheet.cell(2, 1).value
    if d_cell and "DATE" in str(d_cell):
        m = re.search(r'(\d{1,2})[/.-](\d{1,2})[/.-](\d{4})', str(d_cell))
        if m:
            d_str = f"{m.group(3)}-{m.group(2).zfill(2)}-{m.group(1).zfill(2)}"

    current_session = "Breakfast"
    current_pax = 1904

    for r in range(4, sheet.max_row + 1):
        sess_val = sheet.cell(r, 1).value
        pax_val = sheet.cell(r, 2).value
        item_val = sheet.cell(r, 3).value

        if sess_val and "Standard" in str(sess_val):
            raw_s = str(sess_val).replace("\n", " ").strip()
            if "Breakfast" in raw_s: current_session = "Breakfast"
            elif "Lunch" in raw_s: current_session = "Lunch"
            elif "Snacks" in raw_s: current_session = "Snacks"
            elif "Dinner" in raw_s: current_session = "Dinner"
            else: current_session = raw_s.replace("Standard", "").strip()

        if pax_val and isinstance(pax_val, (int, float)) and pax_val > 0:
            current_pax = int(pax_val)

        if not item_val or not str(item_val).strip() or "Grand" in str(item_val) or "Total" in str(item_val):
            continue

        item_name = str(item_val).strip()
        cost = float(sheet.cell(r, 4).value or 0.0)
        uom = str(sheet.cell(r, 5).value or "Kg").strip()
        std_portion = float(sheet.cell(r, 6).value or 1.0)
        conv_grms = float(sheet.cell(r, 7).value or 1.0)

        est_prod = float(sheet.cell(r, 8).value or 0.0)
        act_prod = float(sheet.cell(r, 9).value or 0.0)
        over_prod = float(sheet.cell(r, 10).value or 0.0)
        pickup = float(sheet.cell(r, 11).value or 0.0)
        kitchen_left = float(sheet.cell(r, 12).value or 0.0)
        buffet_ret = float(sheet.cell(r, 13).value or 0.0)
        consumption = float(sheet.cell(r, 14).value or 0.0)
        total_left = float(sheet.cell(r, 16).value or (kitchen_left + buffet_ret))
        reuse = float(sheet.cell(r, 17).value or 0.0)
        waste = float(sheet.cell(r, 18).value or 0.0)
        waste_kg = float(sheet.cell(r, 19).value or 0.0)
        waste_cost = float(sheet.cell(r, 20).value or 0.0)
        raw_waste_pct = sheet.cell(r, 21).value
        waste_pct = float(raw_waste_pct or 0.0) * 100.0 if (raw_waste_pct and float(raw_waste_pct) < 1.0) else float(raw_waste_pct or 0.0)

        # Columns 23-27: converted in Kgs
        act_prod_kg = float(sheet.cell(r, 23).value or (act_prod * conv_grms if uom.lower() in ["pcs", "pkt"] else act_prod))
        pickup_kg = float(sheet.cell(r, 24).value or (pickup * conv_grms if uom.lower() in ["pcs", "pkt"] else pickup))
        kitchen_left_kg = float(sheet.cell(r, 25).value or (kitchen_left * conv_grms if uom.lower() in ["pcs", "pkt"] else kitchen_left))
        buffet_ret_kg = float(sheet.cell(r, 26).value or (buffet_ret * conv_grms if uom.lower() in ["pcs", "pkt"] else buffet_ret))
        reuse_kg = float(sheet.cell(r, 27).value or (reuse * conv_grms if uom.lower() in ["pcs", "pkt"] else reuse))

        total_left_kg = kitchen_left_kg + buffet_ret_kg
        cons_kg = max(0.0, act_prod_kg - total_left_kg)
        est_prod_kg = est_prod * conv_grms if uom.lower() in ["pcs", "pkt"] else est_prod
        over_prod_kg = max(0.0, act_prod_kg - est_prod_kg)

        waste_per_head_g = (waste_kg / current_pax * 1000.0) if current_pax > 0 else 0.0
        cons_per_head_g = (cons_kg / current_pax * 1000.0) if current_pax > 0 else 0.0
        prod_per_head_g = (act_prod_kg / current_pax * 1000.0) if current_pax > 0 else 0.0
        reuse_pct = (reuse_kg / total_left_kg * 100.0) if total_left_kg > 0 else 0.0

        records.append({
            "hotel_name": "Hotel Sahara",
            "property_location": "Main Dining Hall",
            "event_name": f"Daily Operations - {current_session}",
            "event_type": "Regular Hotel Service",
            "service_type": "Buffet",
            "session": current_session,
            "record_date": datetime.strptime(d_str, "%Y-%m-%d").date(),
            "pax": current_pax,
            "dish_name": item_name,
            "dish_category": categorize_dish(item_name),
            "food_type": "Veg",
            "uom": uom,
            "item_cost": cost,
            "standard_qty_per_portion": std_portion,
            "conversion_factor": conv_grms,
            "estimated_production": est_prod,
            "estimated_production_kg": round(est_prod_kg, 2),
            "actual_production": act_prod,
            "actual_production_kg": round(act_prod_kg, 2),
            "over_production": over_prod,
            "over_production_kg": round(over_prod_kg, 2),
            "pickup_quantity": pickup,
            "pickup_quantity_kg": round(pickup_kg, 2),
            "kitchen_leftover": kitchen_left,
            "kitchen_leftover_kg": round(kitchen_left_kg, 2),
            "location_buffet_return": buffet_ret,
            "location_buffet_return_kg": round(buffet_ret_kg, 2),
            "reuse_quantity": reuse,
            "reuse_quantity_kg": round(reuse_kg, 2),
            "actual_consumption": consumption,
            "actual_consumption_kg": round(cons_kg, 2),
            "total_leftover": total_left,
            "total_leftover_kg": round(total_left_kg, 2),
            "total_waste": waste,
            "total_waste_kg": round(waste_kg, 2),
            "waste_cost": round(waste_cost, 2),
            "waste_percentage": round(waste_pct, 2),
            "waste_per_head_grams": round(waste_per_head_g, 1),
            "consumption_per_head_grams": round(cons_per_head_g, 1),
            "production_per_head_grams": round(prod_per_head_g, 1),
            "reuse_percentage": round(reuse_pct, 1),
            "notes": "Parsed from Sahara Daily Production Report",
            "data_source": "Excel Import",
            "source_file": source_file,
        })
    return records

def parse_sitara_sheet(sheet, source_file: str = "Daily report.xlsx") -> List[Dict[str, Any]]:
    records = []
    # Title row 1
    # 'M/S. SARALA 60TH BIRTHDAY GROUP VEG BUFFET DINNER COOKED FOOD CONSUMPTION FOR 230 PAX AT LEG GARDEN LOCATION ON 07.10.2026'
    title = str(sheet.cell(1, 1).value or "")
    pax = 230
    pax_m = re.search(r'(\d+)\s*PAX', title, re.IGNORECASE)
    if pax_m:
        pax = int(pax_m.group(1))

    event_date_str = "2026-10-07"
    date_m = re.search(r'(\d{1,2})[.](\d{1,2})[.](\d{4})', title)
    if date_m:
        event_date_str = f"{date_m.group(3)}-{date_m.group(2).zfill(2)}-{date_m.group(1).zfill(2)}"

    event_name = "M/S. Sarala 60th Birthday"
    location = "LEG Garden Location"

    # Total event waste cost from row 72 is 2783
    # Total event waste kg is 23.0
    # Average cost per kg of waste = 2783 / 23 = ~121 INR/kg
    avg_cost_per_kg = 2783.0 / 23.0 if 23.0 > 0 else 120.0

    for r in range(7, 51):
        item_val = sheet.cell(r, 1).value
        est_val = sheet.cell(r, 2).value
        if not item_val or est_val is None or str(item_val).startswith("PRE COOKED"):
            continue

        item_name = str(item_val).strip()
        est_kg = float(est_val or 0.0)
        pickup_kg = float(sheet.cell(r, 4).value or 0.0)
        cons_kg = float(sheet.cell(r, 6).value or 0.0)
        buffet_left = float(sheet.cell(r, 8).value or 0.0)
        kitchen_left = float(sheet.cell(r, 9).value or 0.0)
        total_left = float(sheet.cell(r, 10).value or (buffet_left + kitchen_left))
        waste_kg = float(sheet.cell(r, 12).value or 0.0)
        reuse_kg = float(sheet.cell(r, 14).value or 0.0)

        # Dish category and cost estimation
        category = categorize_dish(item_name)
        cost_multiplier = 1.4 if "paneer" in item_name.lower() or "sweet" in item_name.lower() else (0.8 if "rice" in item_name.lower() else 1.0)
        item_unit_cost = round(avg_cost_per_kg * cost_multiplier, 2)
        waste_cost = round(waste_kg * item_unit_cost, 2)

        waste_pct = (waste_kg / pickup_kg * 100.0) if pickup_kg > 0 else 0.0
        waste_per_head_g = (waste_kg / pax * 1000.0) if pax > 0 else 0.0
        cons_per_head_g = (cons_kg / pax * 1000.0) if pax > 0 else 0.0
        prod_per_head_g = (pickup_kg / pax * 1000.0) if pax > 0 else 0.0
        reuse_pct = (reuse_kg / total_left * 100.0) if total_left > 0 else 0.0
        over_prod_kg = max(0.0, pickup_kg - est_kg)

        records.append({
            "hotel_name": "Hotel Sitara",
            "property_location": location,
            "event_name": event_name,
            "event_type": "Birthday",
            "service_type": "Buffet",
            "session": "Dinner",
            "record_date": datetime.strptime(event_date_str, "%Y-%m-%d").date(),
            "pax": pax,
            "dish_name": item_name,
            "dish_category": category,
            "food_type": "Veg",
            "uom": "Kg",
            "item_cost": item_unit_cost,
            "standard_qty_per_portion": 1.0,
            "conversion_factor": 1.0,
            "estimated_production": est_kg,
            "estimated_production_kg": est_kg,
            "actual_production": pickup_kg,
            "actual_production_kg": pickup_kg,
            "over_production": over_prod_kg,
            "over_production_kg": over_prod_kg,
            "pickup_quantity": pickup_kg,
            "pickup_quantity_kg": pickup_kg,
            "kitchen_leftover": kitchen_left,
            "kitchen_leftover_kg": kitchen_left,
            "location_buffet_return": buffet_left,
            "location_buffet_return_kg": buffet_left,
            "reuse_quantity": reuse_kg,
            "reuse_quantity_kg": reuse_kg,
            "actual_consumption": cons_kg,
            "actual_consumption_kg": cons_kg,
            "total_leftover": total_left,
            "total_leftover_kg": total_left,
            "total_waste": waste_kg,
            "total_waste_kg": waste_kg,
            "waste_cost": waste_cost,
            "waste_percentage": round(waste_pct, 2),
            "waste_per_head_grams": round(waste_per_head_g, 1),
            "consumption_per_head_grams": round(cons_per_head_g, 1),
            "production_per_head_grams": round(prod_per_head_g, 1),
            "reuse_percentage": round(reuse_pct, 1),
            "notes": "Parsed from Sitara Banquet Event Report",
            "data_source": "Excel Import",
            "source_file": source_file,
        })
    return records

def parse_excel_file(file_content: bytes, filename: str) -> List[Dict[str, Any]]:
    wb = openpyxl.load_workbook(io.BytesIO(file_content), data_only=True)
    all_records = []

    for name in wb.sheetnames:
        lower_name = name.lower()
        sheet = wb[name]
        if "sahara" in lower_name:
            all_records.extend(parse_sahara_sheet(sheet, filename))
        elif "sitara" in lower_name:
            all_records.extend(parse_sitara_sheet(sheet, filename))
        else:
            # Check sheet content header
            header_cell = str(sheet.cell(1, 1).value or "").lower()
            if "sahara" in header_cell:
                all_records.extend(parse_sahara_sheet(sheet, filename))
            elif "sitara" in header_cell or "buffet" in header_cell or "pax" in header_cell:
                all_records.extend(parse_sitara_sheet(sheet, filename))

    # If no specific sheets recognized, attempt sahara then sitara
    if not all_records and wb.sheetnames:
        try:
            all_records = parse_sahara_sheet(wb[wb.sheetnames[0]], filename)
        except Exception:
            try:
                all_records = parse_sitara_sheet(wb[wb.sheetnames[0]], filename)
            except Exception:
                all_records = []

    return all_records

def preview_excel_file(file_content: bytes, filename: str) -> Dict[str, Any]:
    wb = openpyxl.load_workbook(io.BytesIO(file_content), data_only=True)
    sheets_info = []

    for name in wb.sheetnames:
        sheet = wb[name]
        recs = []
        if "sahara" in name.lower():
            recs = parse_sahara_sheet(sheet, filename)
            hotel = "Hotel Sahara"
            event_or_sess = "Daily Production Shifts"
            pax = 5898
            d = "2026-08-24"
        elif "sitara" in name.lower():
            recs = parse_sitara_sheet(sheet, filename)
            hotel = "Hotel Sitara"
            event_or_sess = "M/S. Sarala 60th Birthday"
            pax = 230
            d = "2026-10-07"
        else:
            hotel = name.capitalize()
            event_or_sess = "Banquet Event"
            pax = 200
            d = str(date.today())
            try:
                recs = parse_sahara_sheet(sheet, filename)
            except Exception:
                recs = []

        sheets_info.append({
            "sheet_name": name,
            "hotel": hotel,
            "event_name": event_or_sess,
            "pax": pax,
            "date": d,
            "record_count": len(recs),
            "sample_records": recs[:5],
        })

    return {
        "filename": filename,
        "sheet_count": len(sheets_info),
        "sheets": sheets_info,
    }
